"""
Resume Analyzer Engine:
Extracts text from PDF/DOCX, detects contact information, segments sections,
computes ATS compatibility score, and produces actionable feedback.
"""

import io
import re
from typing import Any, Dict, List, Optional
try:
    import pymupdf as fitz
except ImportError:
    import fitz
import docx

from .skill_extractor import extract_skills, extract_categorized_skills


def extract_text_from_pdf(file_bytes_or_path: Any) -> str:
    """Extracts clean text from a PDF file."""
    text_content = []
    try:
        if isinstance(file_bytes_or_path, (str, bytes, bytearray)):
            doc = fitz.open(stream=file_bytes_or_path, filetype="pdf") if isinstance(file_bytes_or_path, (bytes, bytearray)) else fitz.open(file_bytes_or_path)
        else:
            # Django UploadedFile or file-like object
            stream = file_bytes_or_path.read()
            doc = fitz.open(stream=stream, filetype="pdf")
            if hasattr(file_bytes_or_path, 'seek'):
                file_bytes_or_path.seek(0)

        for page in doc:
            text_content.append(page.get_text())
        doc.close()
    except Exception as e:
        print(f"Error reading PDF: {e}")
    return "\n".join(text_content).strip()


def extract_text_from_docx(file_bytes_or_path: Any) -> str:
    """Extracts text from DOCX paragraphs and tables."""
    text_content = []
    try:
        if isinstance(file_bytes_or_path, str):
            doc = docx.Document(file_bytes_or_path)
        elif isinstance(file_bytes_or_path, (bytes, bytearray)):
            doc = docx.Document(io.BytesIO(file_bytes_or_path))
        else:
            stream = file_bytes_or_path.read()
            doc = docx.Document(io.BytesIO(stream))
            if hasattr(file_bytes_or_path, 'seek'):
                file_bytes_or_path.seek(0)

        for p in doc.paragraphs:
            if p.text.strip():
                text_content.append(p.text.strip())

        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_text:
                    text_content.append(" | ".join(row_text))
    except Exception as e:
        print(f"Error reading DOCX: {e}")
    return "\n".join(text_content).strip()


def extract_text_from_file(file_obj_or_path: Any, filename: str) -> str:
    """Extracts raw text depending on file extension."""
    fn_lower = filename.lower()
    if fn_lower.endswith('.pdf'):
        return extract_text_from_pdf(file_obj_or_path)
    elif fn_lower.endswith(('.docx', '.doc')):
        return extract_text_from_docx(file_obj_or_path)
    elif fn_lower.endswith('.txt'):
        if hasattr(file_obj_or_path, 'read'):
            content = file_obj_or_path.read()
            if hasattr(file_obj_or_path, 'seek'):
                file_obj_or_path.seek(0)
            if isinstance(content, bytes):
                return content.decode('utf-8', errors='ignore')
            return str(content)
        with open(file_obj_or_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
    return ""


def extract_contact_info(text: str) -> Dict[str, Optional[str]]:
    """Extracts email, phone, LinkedIn, GitHub, and portfolio links."""
    # Email
    email_pattern = r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+'
    email_match = re.search(email_pattern, text)
    email = email_match.group(0) if email_match else None

    # Phone (handles +1, +91, (123) 456-7890, 98765 43210, 9876543210, etc.)
    phone_pattern = r'(?:\+?\d{1,3}[-.\s]?)?(?:\(?\d{3,5}\)?[-.\s]?\d{3,5}(?:[-.\s]?\d{3,5})?|\b\d{10}\b)'
    phone_match = re.search(phone_pattern, text)
    phone = phone_match.group(0).strip() if phone_match else None

    # LinkedIn
    linkedin_pattern = r'https?://(?:www\.)?linkedin\.com/in/[a-zA-Z0-9_-]+/?|linkedin\.com/in/[a-zA-Z0-9_-]+/?'
    linkedin_match = re.search(linkedin_pattern, text, re.IGNORECASE)
    linkedin = linkedin_match.group(0) if linkedin_match else None

    # GitHub
    github_pattern = r'https?://(?:www\.)?github\.com/[a-zA-Z0-9_-]+/?|github\.com/[a-zA-Z0-9_-]+/?'
    github_match = re.search(github_pattern, text, re.IGNORECASE)
    github = github_match.group(0) if github_match else None

    # Portfolio / Website
    portfolio_pattern = r'https?://(?!www\.linkedin\.com|www\.github\.com)[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}(?:/[^\s]*)?'
    portfolio_match = re.search(portfolio_pattern, text, re.IGNORECASE)
    portfolio = portfolio_match.group(0) if portfolio_match else None

    return {
        'email': email,
        'phone': phone,
        'linkedin': linkedin,
        'github': github,
        'portfolio': portfolio
    }


def segment_resume_sections(text: str) -> Dict[str, str]:
    """
    Detects and segments standard resume sections:
    Education, Experience, Skills, Projects, Certifications, Summary.
    """
    section_headers = {
        'summary': r'\b(professional summary|summary|objective|profile|about me)\b',
        'education': r'\b(education|academic background|academics|qualifications)\b',
        'experience': r'\b(work experience|experience|employment history|internships|work history)\b',
        'skills': r'\b(technical skills|skills|technologies|key competencies|core skills|tools)\b',
        'projects': r'\b(projects|academic projects|key projects|personal projects)\b',
        'certifications': r'\b(certifications|certificates|licenses|courses|achievements)\b'
    }

    lines = text.split('\n')
    current_section = 'header'
    sections: Dict[str, List[str]] = {
        'header': [],
        'summary': [],
        'education': [],
        'experience': [],
        'skills': [],
        'projects': [],
        'certifications': []
    }

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        # Check if line matches any section header (short line and matches keyword)
        if len(stripped.split()) <= 4:
            matched_sec = None
            for sec_name, pattern in section_headers.items():
                if re.search(pattern, stripped, re.IGNORECASE):
                    matched_sec = sec_name
                    break
            if matched_sec:
                current_section = matched_sec
                continue

        sections[current_section].append(stripped)

    return {k: "\n".join(v).strip() for k, v in sections.items() if v}


def calculate_ats_score(
    text: str,
    contact_info: Dict[str, Optional[str]],
    sections: Dict[str, str],
    skills: List[str]
) -> Dict[str, Any]:
    """
    Calculates ATS compatibility score (0-100) and returns actionable feedback.
    """
    score = 0
    feedback = []

    # 1. Contact Info Completeness (Up to 20 pts)
    contact_pts = 0
    if contact_info.get('email'):
        contact_pts += 6
    else:
        feedback.append({'type': 'warning', 'message': 'Missing email address in resume header.'})

    if contact_info.get('phone'):
        contact_pts += 5
    else:
        feedback.append({'type': 'warning', 'message': 'Missing contact phone number.'})

    if contact_info.get('linkedin'):
        contact_pts += 5
    else:
        feedback.append({'type': 'info', 'message': 'Add your LinkedIn profile URL to boost recruiter visibility.'})

    if contact_info.get('github') or contact_info.get('portfolio'):
        contact_pts += 4
    else:
        feedback.append({'type': 'info', 'message': 'Include a GitHub or portfolio link to showcase your code.'})

    score += contact_pts

    # 2. Key Sections Present (Up to 25 pts)
    sec_pts = 0
    if 'skills' in sections or len(skills) >= 3:
        sec_pts += 7
    else:
        feedback.append({'type': 'danger', 'message': 'Explicit "Technical Skills" section is missing.'})

    if 'experience' in sections:
        sec_pts += 7
    else:
        feedback.append({'type': 'warning', 'message': 'Include an "Experience" or "Internships" section.'})

    if 'education' in sections:
        sec_pts += 6
    else:
        feedback.append({'type': 'warning', 'message': 'Include an "Education" section with degree and year.'})

    if 'projects' in sections:
        sec_pts += 5
    else:
        feedback.append({'type': 'info', 'message': 'Add a "Projects" section demonstrating hands-on implementations.'})

    score += sec_pts

    # 3. Skill Density & Variety (Up to 25 pts)
    skill_pts = 0
    num_skills = len(skills)
    if num_skills >= 10:
        skill_pts = 25
    elif num_skills >= 6:
        skill_pts = 18
    elif num_skills >= 3:
        skill_pts = 12
    else:
        skill_pts = 5
        feedback.append({'type': 'danger', 'message': 'Low keyword density: Add more domain-relevant technologies and tools.'})

    score += skill_pts

    # 4. Action Verbs and Quantifiable Metrics (Up to 15 pts)
    action_verbs = [
        'developed', 'built', 'created', 'implemented', 'designed', 'optimized',
        'improved', 'reduced', 'increased', 'led', 'architected', 'automated',
        'engineered', 'deployed', 'spearheaded', 'delivered'
    ]
    text_lower = text.lower()
    verbs_found = [v for v in action_verbs if v in text_lower]

    metric_pattern = r'\b(\d+%\b|\$\d+|\b\d+x\b|\b\d+\+?\s*(users|clients|requests|ms|seconds|hours|teams|members))'
    metrics_found = re.findall(metric_pattern, text_lower)

    impact_pts = 0
    if len(verbs_found) >= 5:
        impact_pts += 8
    elif len(verbs_found) >= 2:
        impact_pts += 5
    else:
        feedback.append({'type': 'info', 'message': 'Use strong action verbs (e.g., "Architected", "Engineered", "Optimized") at the start of bullet points.'})

    if metrics_found:
        impact_pts += 7
    else:
        feedback.append({'type': 'warning', 'message': 'Add quantifiable metrics and numbers (e.g. "Improved query performance by 40%") to demonstrate impact.'})

    score += impact_pts

    # 5. Length and Formatting Sanity (Up to 15 pts)
    format_pts = 0
    words = text.split()
    word_count = len(words)

    if 300 <= word_count <= 1200:
        format_pts = 15
    elif 150 <= word_count < 300:
        format_pts = 10
        feedback.append({'type': 'info', 'message': 'Resume is somewhat brief. Expand on key projects and responsibilities.'})
    else:
        format_pts = 8
        if word_count > 1200:
            feedback.append({'type': 'info', 'message': 'Resume is lengthy. Try condensing content to 1-2 pages.'})
        else:
            feedback.append({'type': 'danger', 'message': 'Resume content is very minimal. Elaborate on your background.'})

    score += format_pts

    # Clamp score between 0 and 100
    score = min(100, max(0, score))

    if score >= 85:
        feedback.insert(0, {'type': 'success', 'message': 'Excellent ATS formatting! Resume is well-structured and keyword-rich.'})
    elif score >= 70:
        feedback.insert(0, {'type': 'success', 'message': 'Strong resume foundation with good keyword coverage.'})
    else:
        feedback.insert(0, {'type': 'warning', 'message': 'Resume could benefit from optimization for better ATS searchability.'})

    return {
        'score': score,
        'word_count': word_count,
        'action_verbs_count': len(verbs_found),
        'metrics_count': len(metrics_found),
        'feedback': feedback
    }


def analyze_resume(file_obj_or_path: Any, filename: str) -> Dict[str, Any]:
    """
    Main entry point: Analyzes a resume file and returns complete insights.
    """
    raw_text = extract_text_from_file(file_obj_or_path, filename)
    contact_info = extract_contact_info(raw_text)
    sections = segment_resume_sections(raw_text)
    extracted_skills = extract_skills(raw_text)
    categorized_skills = extract_categorized_skills(raw_text)
    ats_data = calculate_ats_score(raw_text, contact_info, sections, extracted_skills)

    return {
        'raw_text': raw_text,
        'contact_info': contact_info,
        'sections': sections,
        'extracted_skills': extracted_skills,
        'categorized_skills': categorized_skills,
        'ats_score': ats_data['score'],
        'ats_feedback': ats_data['feedback'],
        'word_count': ats_data['word_count'],
    }
